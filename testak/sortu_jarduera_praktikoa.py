#!/usr/bin/env python3
"""Ezagutze pasiboa – jarduera praktikoa (domeinu librea + entrega)."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def kontsola(doc: Document, testua: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(0.25)
    pPr = p._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F0F0F0")
    pPr.append(shd)
    run = p.add_run(testua)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)


def galdera(doc: Document, zenbakia: str, testua: str, lerroak: int = 1) -> None:
    p = doc.add_paragraph()
    p.add_run(f"{zenbakia}. ").bold = True
    p.add_run(testua)
    for _ in range(lerroak):
        doc.add_paragraph("_______________________________________________")


def main() -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    h = doc.add_heading("Ezagutze pasiboa – Jarduera praktikoa", level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    azpi = doc.add_paragraph()
    azpi.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = azpi.add_run(
        "whois · nslookup · dig · hosting mota (interneten bilatu)\n"
        "Ikasleak aukeratutako domeinu batekin"
    )
    r.italic = True

    doc.add_heading("Helburua", level=1)
    doc.add_paragraph(
        "Domeinu publiko bat aukeratu eta ezagutze pasiboa egin "
        "whois, nslookup eta dig tresnekin. Amaieran, zure bilaketen "
        "emaitzak dokumentu batean entregatu."
    )

    doc.add_heading("Entrega", level=1)
    doc.add_paragraph(
        "Entregatu Word dokumentu bat (.docx). "
        "Izenburuaren adibidea: Ezagutze_pasiboa_IzenaAbizena.docx. "
        "Dokumentuak honakoak izan behar ditu:"
    )
    for item in (
        "Aukeratutako domeinua eta zergatik aukeratu duzun (1–2 esaldi).",
        "Exekutatutako komando bakoitza (kopiatu-itsatsi).",
        "Komando bakoitzaren irteeraren zati garrantzitsua (ez derrigorrez dena).",
        "Beheko galderen erantzunak.",
        "Hosting / zerbitzari motari buruzko bilaketa (interneteko iturriak aipatuz).",
        "Ondorio labur bat (zer ikasi duzun / zer aurkitu duzun).",
    ):
        doc.add_paragraph(item, style="List Number")

    doc.add_heading("Arauak", level=1)
    for item in (
        "Domeinu LIBREA aukeratu hasieratik (webgune, enpresa, eskola, proiektu…).",
        "Ez erabili TryHackMe-ko adibide finkoak (tryhackme.com, thmlabs.com) lanaren oinarri gisa.",
        "Soilik kontsulta pasiboak: whois, nslookup, dig. Ez ping, ez nmap, ez eskaneorik.",
        "Errespetatu pribatutasuna: ez bilatu datu pertsonal sentikorrak entregatzeko.",
        "Taldean egiten bada, adierazi kideak dokumentuan.",
    ):
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("0. Prestaketa", level=1)
    doc.add_paragraph("Tresnak egiaztatu:")
    kontsola(
        doc,
        "which whois nslookup dig\n"
        "# Faltaren bat badago (Debian/Ubuntu):\n"
        "sudo apt update && sudo apt install -y whois dnsutils",
    )
    doc.add_paragraph(
        "Aukeratu zure DOMEINUA (adib. wikipedia.org, cloudflare.com, "
        "zure herriko udalaren webgunea…). Ordezkatu DOMEINUA zure aukerarekin "
        "hurrengo komando guztietan."
    )
    galdera(doc, "0.1", "Zure domeinua:")
    galdera(doc, "0.2", "Zergatik aukeratu duzu? (1–2 esaldi)", lerroak=2)

    # --- whois ---
    doc.add_heading("1. whois", level=1)
    doc.add_paragraph(
        "WHOIS-ek domeinuaren erregistro-datu publikoak ematen ditu "
        "(erregistratzailea, datak, name-zerbitzariak…)."
    )
    doc.add_paragraph("Exekutatu eta gorde irteera (edo zati garrantzitsua):")
    kontsola(doc, "whois DOMEINUA")
    doc.add_paragraph("Luzeegia bada:")
    kontsola(doc, "whois DOMEINUA | less")

    galdera(doc, "1.1", "Creation / Registration Date (erregistro-data):")
    galdera(doc, "1.2", "Registrar (erregistratzailea):")
    galdera(doc, "1.3", "Name Server-ak (gutxienez bi, badaude):", lerroak=2)
    galdera(
        doc,
        "1.4",
        "WHOIS irteeran, zer informazio dago redaktatuta / ezkutatuta pribatutasunagatik? "
        "Ez badago ezer ezkutatuta, idatzi “ezer ez”.",
        lerroak=2,
    )
    galdera(
        doc,
        "1.5",
        "Galderatxoa: WHOIS zerbitzariek tradizionalki zein TCP atakan entzuten dute?",
    )

    # --- nslookup ---
    doc.add_heading("2. nslookup", level=1)
    doc.add_paragraph(
        "nslookup-ek DNS erregistroak kontsultatzen ditu. "
        "Erabilgarria da, batez ere Windows inguruneetan eta dokumentazio zaharrean."
    )
    doc.add_paragraph("A (IPv4), MX eta TXT:")
    kontsola(
        doc,
        "nslookup DOMEINUA\n"
        "nslookup -type=MX DOMEINUA\n"
        "nslookup -type=TXT DOMEINUA\n"
        "# Aukerakoa: DNS zerbitzari publikoa zehaztu\n"
        "nslookup -type=MX DOMEINUA 1.1.1.1",
    )

    galdera(doc, "2.1", "A erregistroa: lortutako IPv4 helbidea(k):", lerroak=2)
    galdera(doc, "2.2", "MX: posta-zerbitzariak eta priority balioak (badaude):", lerroak=2)
    galdera(
        doc,
        "2.3",
        "TXT: aipatu aurkitutako sarrera interesgarri bat (edo “ez dago / hutsa”):",
        lerroak=2,
    )
    galdera(
        doc,
        "2.4",
        "Galderatxoa: nslookup -type=AAAA DOMEINUA eginez, zer motatako helbideak bilatzen dituzu?",
    )

    # --- dig ---
    doc.add_heading("3. dig", level=1)
    doc.add_paragraph(
        "dig tresna modernoa da: irteera argiagoa, TTL balioak erakusten ditu "
        "eta scripting-erako egokiagoa da."
    )
    kontsola(
        doc,
        "dig DOMEINUA A\n"
        "dig DOMEINUA MX\n"
        "dig DOMEINUA TXT\n"
        "dig DOMEINUA NS\n"
        "# DNS zerbitzari zehatza:\n"
        "dig @1.1.1.1 DOMEINUA MX",
    )

    galdera(
        doc,
        "3.1",
        "A erregistroaren TTL balioa (ANSWER SECTION-eko segundoak):",
    )
    galdera(doc, "3.2", "NS erregistroak (name-zerbitzariak dig-ekin):", lerroak=2)
    galdera(
        doc,
        "3.3",
        "nslookup-ekin lortutako A / MX emaitzak dig-ekin bat datoz? "
        "Azaldu labur desberdintasunik badago (irteeraren formatua, TTL…).",
        lerroak=2,
    )
    galdera(
        doc,
        "3.4",
        "Galderatxoa: dig @8.8.8.8 DOMEINUA A komandoan, zer adierazten du @8.8.8.8 zatiak?",
    )

    # --- Hosting / zerbitzari mota (interneten bilatu) ---
    doc.add_heading("4. Zerbitzari / hosting mota (interneten bilatu)", level=1)
    doc.add_paragraph(
        "Webguneak hosting mota desberdinetan egon daitezke, adibidez:"
    )
    for item in (
        "Shared hosting (ostatze partekatua) – makina bera bezero askoren artean.",
        "VPS (Virtual Private Server) – makina birtuala, baliabide propioagoak.",
        "Dedicated – zerbitzari fisiko osoa bezero bakar batentzat.",
        "Cloud / CDN – hodeiko azpiegitura edo edukia banatzeko sarea "
        "(adib. Cloudflare, AWS, Azure…).",
    ):
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph(
        "Zure domeinuari buruzko pistak erabili (WHOIS-eko registrar, name-zerbitzariak, "
        "IP-aren jabea…) eta bilatu Interneten hosting / zerbitzari mota. "
        "Adibideak: hosting-en dokumentazioa, “who hosts”, IP/ASN bilatzaileak, "
        "enpresaren “about / hosting” orriak. Aipatu erabilitako iturria."
    )
    doc.add_paragraph(
        "Oharra: batzuetan ez da %100 ziurra; azaldu zure hipotesia eta zertan oinarritu zaren."
    )

    galdera(
        doc,
        "4.1",
        "Zure ustez, zer motatako hosting / zerbitzaria da? "
        "(shared, VPS, dedicated, cloud/CDN, beste…)",
    )
    galdera(
        doc,
        "4.2",
        "Zertan oinarritu zara? (NS izenak, IP/hornitzailea, web bilaketa…)",
        lerroak=2,
    )
    galdera(
        doc,
        "4.3",
        "Erabilitako Interneteko iturria(k) (URL edo tresnaren izena):",
        lerroak=2,
    )
    galdera(
        doc,
        "4.4",
        "Galderatxoa: shared hosting batean, zergatik izan daiteke arriskutsuagoa "
        "edo interesgarriagoa erasotzaile batentzat (ideiaz, labur)?",
        lerroak=2,
    )

    # --- Hiruren inguruko galderak ---
    doc.add_heading("5. Hiru tresnen inguruko galderatxoak", level=1)
    galdera(
        doc,
        "5.1",
        "Zergatik jotzen dira whois, nslookup eta dig kontsulta hauek ezagutze pasibotzat?",
        lerroak=2,
    )
    galdera(
        doc,
        "5.2",
        "Zein tresnak ematen ditu domeinuaren erregistro-datuak (data, registrar…), "
        "eta zeinek DNS erregistroak (A, MX, TXT…)?",
        lerroak=2,
    )
    galdera(
        doc,
        "5.3",
        "nslookup eta dig artean, zein gomendatuko zenuke gaur egun eta zergatik?",
        lerroak=2,
    )
    galdera(
        doc,
        "5.4",
        "Adibide bat: zer litzateke ezagutze AKTIBOA egoera berean? "
        "(ez egin; azaldu soilik)",
        lerroak=2,
    )

    # --- Ondorioa ---
    doc.add_heading("6. Ondorioa", level=1)
    doc.add_paragraph(
        "Idatzi 4–6 esalditan: zer aurkitu duzun zure domeinuari buruz "
        "(tresnak + hosting mota), zer izan den erabilgarriena, eta zer kontuz "
        "ibili behar den informazio publikoa bilatzean."
    )
    for _ in range(5):
        doc.add_paragraph("_______________________________________________")

    doc.add_paragraph()
    ohar = doc.add_paragraph()
    ohar.add_run("Ebaluazio-iradokizuna (irakaslea): ").bold = True
    ohar.add_run(
        "domeinu propioa · hiru tresnak · hosting mota + iturria · "
        "komandoak + irteerak · galderak · ondorioa."
    )

    # --- Irakaslearen gida ---
    doc.add_page_break()
    gh = doc.add_heading("Irakaslearen gida (ez banatu ikasleekin)", level=1)
    gh.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        "Ikasle bakoitzak (edo taldeak) domeinu desberdina aukeratuko du; "
        "erantzun faktikoak (IP, TTL, registrar…) ez dira bakarrak. "
        "Egiaztatu prozesua eta ulermena."
    )

    gida = [
        ("1.5", "43 (TCP)."),
        ("2.4", "IPv6 helbideak (AAAA erregistroak)."),
        (
            "3.4",
            "Kontsulta 8.8.8.8 DNS zerbitzariari (Google DNS) bidaltzen zaiola.",
        ),
        (
            "4.1–4.3",
            "Ikaslearen arabera (shared / VPS / dedicated / cloud-CDN…). "
            "Egiaztatu hipotesia + iturria aipatu dituen. Cloudflare NS = CDN/proxy "
            "pista ohikoa; ez nahastu “registrar” eta “hosting”.",
        ),
        (
            "4.4",
            "Baliabideak partekatzen dira: auzoko bezero baten ahultasunak "
            "eragina izan dezake; gainera, askotan konfigurazio eta isolamendu ahulagoa.",
        ),
        (
            "5.1",
            "Iturri publikoak kontsultatzen dira; ez zaio paketerik bidaltzen "
            "helburuko ostalariari zuzenean.",
        ),
        (
            "5.2",
            "WHOIS → erregistro-datuak. nslookup eta dig → DNS erregistroak.",
        ),
        (
            "5.3",
            "dig: irteera argiagoa, TTL, scripting. nslookup: bateragarritasuna.",
        ),
        (
            "5.4",
            "Adib.: nmap, ping, direktorio-brute force, helburuarekin interakzio zuzena…",
        ),
    ]
    for zb, t in gida:
        p = doc.add_paragraph()
        p.add_run(f"{zb}. ").bold = True
        p.add_run(t)

    irteerak = [
        Path("/workspace/testak/Ezagutze_pasiboa_jarduera_praktikoa.docx"),
        Path("/opt/cursor/artifacts/Ezagutze_pasiboa_jarduera_praktikoa.docx"),
    ]
    for bidea in irteerak:
        bidea.parent.mkdir(parents=True, exist_ok=True)
        doc.save(bidea)
        print(f"Gordeta: {bidea}")


if __name__ == "__main__":
    main()

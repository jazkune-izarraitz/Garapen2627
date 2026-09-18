#!/usr/bin/env python3
"""Ezagutze pasiboa moduluko 15 galderako testa (Word)."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# Galdera bakoitzeko aukerak: (testua, mota) — mota: zuzena | txikia | handia
# "kontsola" = galdera praktikoa (TryHackMe estiloa). Guztira 5 praktiko.
GALDERAK = [
    {
        "zenbakia": 1,
        "galdera": "Zer da ezagutze pasiboa (passive reconnaissance)?",
        "aukerak": [
            ("Helburuko sisteman sartzea eta fitxategiak aldatzea.", "handia"),
            (
                "Helburuaren informazioa iturri publikoetatik biltzea, helburuarekin zuzenean kontaktatu gabe.",
                "zuzena",
            ),
            (
                "Helburuaren informazioa iturri publikoetatik biltzea, baina Nmap erabiliz soilik.",
                "txikia",
            ),
        ],
    },
    {
        "zenbakia": 2,
        "galdera": "Enpresa baten Facebook orria bisitatzen duzu langileen izenak lortzeko. Zer motatako ezagutzea da?",
        "aukerak": [
            ("Pasiboa, informazio publikoa soilik erabiltzen delako.", "zuzena"),
            ("Ez da ezagutzea; marketing jarduera bat da.", "handia"),
            (
                "Aktiboa, sare sozialak beti ezagutze aktibotzat hartzen direlako.",
                "txikia",
            ),
        ],
    },
    {
        "zenbakia": 3,
        "galdera": "Enpresaren web-zerbitzariaren IP helbidea ping egiten duzu ICMP blokeatuta dagoen jakiteko. Zer motatako ezagutzea da?",
        "aukerak": [
            ("Pasiboa, ping komandoa DNS kontsulta bat delako.", "txikia"),
            ("Ez da ezagutzea; sarearen mantentze-lan arrunta da.", "handia"),
            ("Aktiboa, helburuari paketeak bidaltzen zaizkiolako.", "zuzena"),
        ],
    },
    # PRAKTIKOA 1/5
    {
        "zenbakia": 4,
        "galdera": "WHOIS irteera honetan, noiz erregistratu zen tryhackme.com?",
        "kontsola": """ubuntu@tryhackme$ whois tryhackme.com
[Querying whois.verisign-grs.com]
[Redirected to whois.namecheap.com]
Domain Name: TRYHACKME.COM
Registrar: NAMECHEAP INC
Creation Date: 2018-07-05T19:46:15.00Z
Updated Date: 2021-05-01T19:43:23.31Z
Registrar Registration Expiration Date: 2027-07-05T19:46:15.00Z
Name Server: NS1.CLOUDFLARE.COM
Name Server: NS2.CLOUDFLARE.COM""",
        "aukerak": [
            ("2021-05-01", "txikia"),
            ("2018-07-05", "zuzena"),
            ("2027-07-05", "handia"),
        ],
    },
    {
        "zenbakia": 5,
        "galdera": "Tradizionalki, WHOIS zerbitzariek zein TCP atakan entzuten dute?",
        "aukerak": [
            ("53", "txikia"),
            ("80", "handia"),
            ("43", "zuzena"),
        ],
    },
    # PRAKTIKOA 2/5
    {
        "zenbakia": 6,
        "galdera": "dig irteera honetan, zein da thmlabs.com domeinuaren TXT erregistroaren balioa?",
        "kontsola": """ubuntu@tryhackme$ dig thmlabs.com TXT

; <<>> DiG 9.16.1-Ubuntu <<>> thmlabs.com TXT
;; QUESTION SECTION:
;thmlabs.com.                   IN      TXT

;; ANSWER SECTION:
thmlabs.com.            300     IN      TXT     "THM{a5b83929888ed36acb0272971e438d78}"

;; Query time: 28 msec
;; SERVER: 1.1.1.1#53(1.1.1.1)""",
        "aukerak": [
            ('"THM{a5b83929888ed36acb0272971e438d78}"', "zuzena"),
            ("300 (TTL balioa da bandera)", "txikia"),
            ("1.1.1.1 (DNS zerbitzariaren IP-a da bandera)", "handia"),
        ],
    },
    # PRAKTIKOA 3/5
    {
        "zenbakia": 7,
        "galdera": "Komando hau exekutatzen duzu. Zertarako balio du zehazki?",
        "kontsola": """ubuntu@tryhackme$ dig @1.1.1.1 tryhackme.com MX

; <<>> DiG 9.16.1-Ubuntu <<>> @1.1.1.1 tryhackme.com MX
;; QUESTION SECTION:
;tryhackme.com.                 IN      MX

;; ANSWER SECTION:
tryhackme.com.          300     IN      MX      10 alt1.aspmx.l.google.com.
tryhackme.com.          300     IN      MX      20 alt2.aspmx.l.google.com.
tryhackme.com.          300     IN      MX      30 aspmx.l.google.com.""",
        "aukerak": [
            (
                "tryhackme.com-en posta-zerbitzariak (MX) kontsultatzea Cloudflare DNS (1.1.1.1) erabiliz.",
                "zuzena",
            ),
            (
                "tryhackme.com-en A (IPv4) erregistroak kontsultatzea 1.1.1.1 zerbitzarian.",
                "txikia",
            ),
            (
                "tryhackme.com-en WHOIS datuak lortzea TCP 43 atakan.",
                "handia",
            ),
        ],
    },
    {
        "zenbakia": 8,
        "galdera": "DNSDumpster bezalako tresnek zertarako laguntzen dute ezagutze pasiboan?",
        "aukerak": [
            (
                "Azpidomeinuak aurkitzeko, baina soilik helburuari Nmap bidalita.",
                "txikia",
            ),
            (
                "Azpidomeinuak, DNS erregistroak eta azpiegituraren mapa bisuala lortzeko.",
                "zuzena",
            ),
            ("Malwarea instalatzeko eta saioak bahitzeko.", "handia"),
        ],
    },
    {
        "zenbakia": 9,
        "galdera": "Shodan.io zer da?",
        "aukerak": [
            ("Antivirus programa bat ordenagailu pertsonaletarako.", "handia"),
            (
                "Internetera konektatutako gailu eta zerbitzuen bilatzailea (OSINT).",
                "zuzena",
            ),
            ("Soilki DNS MX erregistroak bilatzeko tresna.", "txikia"),
        ],
    },
    {
        "zenbakia": 10,
        "galdera": "Zer esan nahi du DNS siglak?",
        "aukerak": [
            ("Domain Network Service", "txikia"),
            ("Data Numbering Software", "handia"),
            ("Domain Name System", "zuzena"),
        ],
    },
    # PRAKTIKOA 4/5
    {
        "zenbakia": 11,
        "galdera": "nslookup irteera honetan, zein erregistro mota agertzen da eta zer IP mota da?",
        "kontsola": """ubuntu@tryhackme$ nslookup -type=AAAA website.thm

Server:         1.1.1.1
Address:        1.1.1.1#53

Non-authoritative answer:
Name:   website.thm
Address: 2606:4700:20::681a:be5""",
        "aukerak": [
            ("AAAA erregistroa → IPv6 helbidea", "zuzena"),
            ("A erregistroa → IPv6 helbidea ere onartzen duelako", "txikia"),
            ("MX erregistroa → posta-zerbitzariaren IPv6-a", "handia"),
        ],
    },
    {
        "zenbakia": 12,
        "galdera": "Zein erregistro motak adierazten du nora bidali behar den posta elektronikoa?",
        "aukerak": [
            ("CNAME", "txikia"),
            ("MX", "zuzena"),
            ("AAAA", "handia"),
        ],
    },
    {
        "zenbakia": 13,
        "galdera": "CNAME erregistro batek zer egiten du?",
        "aukerak": [
            ("Posta-zerbitzarien lehentasuna (priority) ezartzen du.", "handia"),
            ("Hostname bat IPv4 helbide batera soilik lotzen du.", "txikia"),
            (
                "Hostname bat beste domeinu-izen batera (alias) lotzen du.",
                "zuzena",
            ),
        ],
    },
    # PRAKTIKOA 5/5
    {
        "zenbakia": 14,
        "galdera": "dig irteera honetan, zenbat segundo gorde daiteke erregistroa cachean (TTL)?",
        "kontsola": """ubuntu@tryhackme$ dig www.website.thm A

;; QUESTION SECTION:
;www.website.thm.               IN      A

;; ANSWER SECTION:
www.website.thm.        3600    IN      A       10.10.10.10""",
        "aukerak": [
            ("3600 segundo", "zuzena"),
            ("10 segundo (IP-aren azken oktetoa)", "txikia"),
            ("53 segundo (DNS ataka)", "handia"),
        ],
    },
    {
        "zenbakia": 15,
        "galdera": "Zein DNS zerbitzari motak gordetzen ditu domeinu baten erregistro guztiak?",
        "aukerak": [
            ("WHOIS zerbitzaria, TCP 43 atakan", "handia"),
            ("Authoritative (agintaritza-zerbitzaria)", "zuzena"),
            (
                "Recursive (ISP-k ematen duena), erregistro guztiak beti bertan daudelako",
                "txikia",
            ),
        ],
    },
]

LETRAK = ("a", "b", "c")
PRAKTIKOAK = [g["zenbakia"] for g in GALDERAK if g.get("kontsola")]


def aukera_letra(galdera: dict, mota: str) -> str:
    for i, (_testua, m) in enumerate(galdera["aukerak"]):
        if m == mota:
            return LETRAK[i]
    raise ValueError(f"Mota ez da aurkitu: {mota}")


def gehitu_kontsola(doc: Document, testua: str) -> None:
    """Kontsolako irteera monospace + atzeko gris arinarekin."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Cm(0.25)

    pPr = p._element.get_or_add_pPr()
    shd_el = OxmlElement("w:shd")
    shd_el.set(qn("w:val"), "clear")
    shd_el.set(qn("w:color"), "auto")
    shd_el.set(qn("w:fill"), "F0F0F0")
    pPr.append(shd_el)

    run = p.add_run(testua)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)


def main() -> None:
    assert len(PRAKTIKOAK) == 5, f"5 praktiko behar; daude: {PRAKTIKOAK}"

    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    izenburua = doc.add_heading("Ezagutze pasiboa – Testa", level=0)
    izenburua.alignment = WD_ALIGN_PARAGRAPH.CENTER

    azpi = doc.add_paragraph()
    azpi.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = azpi.add_run(
        "TryHackMe: Passive Reconnaissance eta DNS in Detail\n"
        "15 galdera · 3 aukera (aukera bakarra zuzena da)\n"
        f"5 galdera praktikoak (kontsola): {', '.join(str(n) for n in PRAKTIKOAK)}"
    )
    run.italic = True
    run.font.size = Pt(11)

    info = doc.add_paragraph()
    info.add_run(
        "Jarraibideak: Galdera bakoitzean aukeratu erantzun zuzena (a, b edo c). "
        "Galdera praktikoetan, irakurri kontsolako irteera eta erantzun horren arabera. "
        "Erantzunak dokumentuko hurrengo orrialdean daude."
    )

    doc.add_paragraph()

    for g in GALDERAK:
        p = doc.add_paragraph()
        etiketa = " [praktikoa]" if g.get("kontsola") else ""
        p.add_run(f"{g['zenbakia']}. {g['galdera']}{etiketa}").bold = True
        if g.get("kontsola"):
            gehitu_kontsola(doc, g["kontsola"])
        for letra, (testua, _mota) in zip(LETRAK, g["aukerak"]):
            doc.add_paragraph(f"{letra}) {testua}")
        doc.add_paragraph()

    doc.add_page_break()

    erantzun_izenburua = doc.add_heading("Erantzunak", level=1)
    erantzun_izenburua.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        "Galdera bakoitzean: erantzun zuzena, akats txikia eta akats handia azaltzen dira. "
        f"Galdera praktikoak (kontsola): {', '.join(str(n) for n in PRAKTIKOAK)}."
    )
    doc.add_paragraph()

    for g in GALDERAK:
        zuzena_letra = aukera_letra(g, "zuzena")
        txikia_letra = aukera_letra(g, "txikia")
        handia_letra = aukera_letra(g, "handia")

        zuzena_testua = next(t for t, m in g["aukerak"] if m == "zuzena")
        txikia_testua = next(t for t, m in g["aukerak"] if m == "txikia")
        handia_testua = next(t for t, m in g["aukerak"] if m == "handia")

        p = doc.add_paragraph()
        p.add_run(f"{g['zenbakia']}. ").bold = True
        etiketa = " [praktikoa] " if g.get("kontsola") else " "
        p.add_run(f"{etiketa}Zuzena: {zuzena_letra}) {zuzena_testua}")

        doc.add_paragraph(
            f"   · Akats txikia: {txikia_letra}) {txikia_testua}",
            style="List Bullet",
        )
        doc.add_paragraph(
            f"   · Akats handia: {handia_letra}) {handia_testua}",
            style="List Bullet",
        )
        doc.add_paragraph()

    irteerak = [
        Path("/workspace/testak/Ezagutze_pasiboa_testa.docx"),
        Path("/opt/cursor/artifacts/Ezagutze_pasiboa_testa.docx"),
    ]
    for bidea in irteerak:
        bidea.parent.mkdir(parents=True, exist_ok=True)
        doc.save(bidea)
        print(f"Gordeta: {bidea}")
    print(f"Praktikoak: {PRAKTIKOAK}")


if __name__ == "__main__":
    main()
